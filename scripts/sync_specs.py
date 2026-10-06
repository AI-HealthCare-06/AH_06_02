# /// script
# requires-python = ">=3.13"
# dependencies = ["openpyxl>=3.1"]
# ///
"""구글 시트 사본을 docs/01_planning 아래 마크다운으로 다시 뽑는다.

쓰는 법
    1. 구글 시트에서 파일 → 다운로드 → Microsoft Excel(.xlsx) 로 세 개를 받는다.
       당고킬러_테이블명세서 · 당고킬러_요구사항정의서 · 당고킬러_API 명세서
    2. uv run scripts/sync_specs.py
       기본값은 ~/Downloads 에서 가장 최근 파일을 찾는다. 이름에 (1) (2) 가 붙어도 된다.
       파일 이름 앞부분은 그대로 두어야 한다. 바꾸면 못 찾는다.
    3. git diff 로 바뀐 데만 확인하고 커밋한다.

    다른 폴더에 받았으면  uv run scripts/sync_specs.py --src ~/바탕화면
    한 문서만 다시 뽑으려면 uv run scripts/sync_specs.py --only table

시트가 원본이다. 이 스크립트가 만든 마크다운을 손으로 고치지 않는다.
"""

from __future__ import annotations

import argparse
import datetime as dt
import re
import sys
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

REPO_ROOT = Path(__file__).resolve().parent.parent
OUT_DIR = REPO_ROOT / "docs" / "01_planning"

# 파일 이름이 이 조각으로 시작해야 해당 문서로 본다.
# '테이블' 처럼 느슨하게 잡으면 Downloads 에 있는 남의 자료나 옛날 초안이 걸린다.
SOURCES = {
    "table": ("당고킬러_테이블명세서", "table-spec.md", "당고킬러_테이블명세서"),
    "req": ("당고킬러_요구사항정의서", "requirements.md", "당고킬러_요구사항정의서"),
    "api": ("당고킬러_API", "api-spec.md", "당고킬러_API 명세서"),
}


# ---------------------------------------------------------------- 공통


def cell(value: Any) -> str:
    """셀 값을 마크다운 표 한 칸에 넣을 수 있는 한 줄 문자열로 만든다."""
    if value is None:
        return ""
    if isinstance(value, dt.datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).strip()
    text = text.replace("\r\n", " ").replace("\n", " ").replace("|", "\\|")
    return re.sub(r"\s{2,}", " ", text)


def rows_of(ws: Worksheet) -> list[list[str]]:
    return [[cell(c) for c in row] for row in ws.iter_rows(values_only=True)]


def is_blank(row: list[str]) -> bool:
    return not any(row)


def table(header: list[str], body: list[list[str]], width: int | None = None) -> list[str]:
    width = width or len(header)
    out = ["| " + " | ".join(pad(header, width)) + " |", "| " + " | ".join(["---"] * width) + " |"]
    out += ["| " + " | ".join(pad(r, width)) + " |" for r in body]
    return out


def pad(row: list[str], width: int) -> list[str]:
    return (row + [""] * width)[:width]


def find_sheet(wb: Any, *needles: str) -> Worksheet | None:
    for name in wb.sheetnames:
        flat = name.replace(" ", "")
        if any(n.replace(" ", "") in flat for n in needles):
            return wb[name]
    return None


def header_block(title: str, source: str, what: str, today: str) -> list[str]:
    return [
        f"# {title}",
        "",
        f"> 원본은 구글 시트 `{source}`입니다. 이 파일은 {today} 기준 사본입니다.",
        f"> {what} 바꿀 때는 시트를 먼저 고치고 팀에 알린 뒤 이 파일을 다시 뽑습니다.",
        "> `uv run scripts/sync_specs.py` 로 만듭니다. 손으로 고치지 마세요.",
        "",
    ]


# ---------------------------------------------------------------- 테이블 명세서


def _table_list_section(wb: Any) -> tuple[list[str], dict[str, str]]:
    """테이블 목록 탭. 컬럼 수는 아래 컬럼 명세의 소제목에서 다시 쓴다."""
    ws = find_sheet(wb, "테이블목록")
    if ws is None:
        return [], {}
    rows = [r for r in rows_of(ws) if not is_blank(r)]
    out = ["## 테이블 목록", ""]
    out += table(pad(rows[0], 7), [pad(r, 7) for r in rows[1:]], 7)
    out += [""]
    counts = {r[1]: r[5] for r in rows[1:] if len(r) > 5 and r[1]}
    return out, counts


def _columns_section(wb: Any, counts: dict[str, str]) -> list[str]:
    """테이블 명세 탭. A열에 테이블명이 있는 행에서 새 테이블이 시작된다."""
    ws = find_sheet(wb, "테이블명세")
    if ws is None:
        return []
    out = ["## 테이블별 컬럼 명세", ""]
    started = False
    for row in rows_of(ws)[1:]:
        name = row[0] if row else ""
        if name:
            if started:
                out += [""]
            out += [f"### {name} — {pad(row, 3)[1]}"]
            out += [f"담당 {pad(row, 3)[2]} · {counts.get(name, '?')}컬럼", ""]
            out += table(["No", "컬럼명", "타입", "NULL", "키", "기본값", "설명", "상태"], [], 8)
            started = True
        if len(row) > 4 and row[4]:
            out += ["| " + " | ".join(pad(row, 11)[3:11]) + " |"]
    return out + [""]


def _relations_section(wb: Any) -> list[str]:
    ws = find_sheet(wb, "관계")
    if ws is None:
        return []
    rows = [r for r in rows_of(ws) if not is_blank(r)]
    out = ["## 관계 (FK)", ""]
    out += table(["No", "FROM", "TO", "관계", "설명", "합의 필요"], [pad(r, 6) for r in rows[1:]], 6)
    return out + [""]


def _rules_section(wb: Any) -> list[str]:
    ws = find_sheet(wb, "규칙과범례", "규칙")
    if ws is None:
        return []
    rows = [r for r in rows_of(ws) if not is_blank(r)]
    out = ["## 규칙과 범례", ""]
    out += table(["항목", "내용"], [pad(r, 2) for r in rows], 2)
    return out + [""]


def build_table_spec(wb: Any, today: str) -> str:
    out = header_block("당고킬러 테이블 명세서", "당고킬러_테이블명세서", "스키마를", today)
    listing, counts = _table_list_section(wb)
    out += listing
    out += _columns_section(wb, counts)
    out += _relations_section(wb)
    out += _rules_section(wb)
    return "\n".join(out)


# ---------------------------------------------------------------- 요구사항 정의서


def build_requirements(wb: Any, today: str) -> str:
    ws = find_sheet(wb, "요구사항정의서") or wb[wb.sheetnames[0]]
    rows = [r for r in rows_of(ws)[1:] if not is_blank(r) and pad(r, 1)[0]]

    functional = sum(1 for r in rows if pad(r, 7)[6] == "기능")
    out = header_block("당고킬러 요구사항 정의서", "당고킬러_요구사항정의서", "요구사항을", today)
    out += [f"총 {len(rows)}항목 · 기능 {functional} · 비기능 {len(rows) - functional}", ""]

    current = None
    for row in rows:
        rid, category, d1, d2, d3, body, kind, priority, note = pad(row, 9)
        if category != current:
            out += ["", f"## {category}", ""]
            current = category
        depth = " · ".join(x for x in (d1, d2, d3) if x)
        out += [f"### {rid} — {depth}" if depth else f"### {rid}", ""]
        out += [body, ""]
        out += [f"*{kind} · 우선순위 {priority}*", ""]
        if note:
            out += [f"> 비고: {note}", ""]

    return "\n".join(out).replace("\n\n\n\n", "\n\n")


# ---------------------------------------------------------------- API 명세서
#
# 모든 탭이 같은 모양이다.
#   1행 제목          "API 명세 — A · 회원·인증"
#   2행 부제          "담당 배수빈 · 11개 · 작성 완료 2026-09-28"
#   3행 헤더          API ID | 기능 | Method | Endpoint | 인증 | ... | 상태
#   4행부터 내용
# 헤더 첫 칸이 API ID면 엔드포인트 탭으로 보고 항목별로 풀어 쓴다.

COUNT_RE = re.compile(r"·?\s*(\d+개)\s*·?")


def _api_heading(title: str, subtitle: str) -> tuple[str, str]:
    """제목에서 'API 명세 —' 를 떼고, 부제에 있는 'N개'를 제목 쪽으로 옮긴다."""
    heading = re.sub(r"^API\s*명세\s*[—-]\s*", "", title).strip()
    found = COUNT_RE.search(subtitle)
    if found:
        heading = f"{heading} · {found.group(1)}"
        subtitle = COUNT_RE.sub(" · ", subtitle)
        subtitle = re.sub(r"\s{2,}", " ", subtitle).replace("· ·", "·").strip(" ·")
    return heading, subtitle


def _api_endpoints(head: list[str], body: list[list[str]]) -> list[str]:
    width = len(head)
    out: list[str] = []
    for row in body:
        r = pad(row, width)
        if not r[0]:
            continue
        out += [f"### {r[0]} · {r[1]}", ""]
        out += [f"`{r[2]} {r[3]}`".strip().replace("` ", "`", 1) if r[2] or r[3] else "", ""]
        for i in range(4, width):
            if not head[i]:
                continue
            out += [f"- **{head[i]}**: {r[i] or '—'}"]
        out += [""]
    return out


def build_api_spec(wb: Any, today: str) -> str:
    out = header_block("당고킬러 API 명세서", "당고킬러_API 명세서", "계약을", today)

    for name in wb.sheetnames:
        rows = [r for r in rows_of(wb[name]) if not is_blank(r)]
        if len(rows) < 3:
            continue

        heading, subtitle = _api_heading(pad(rows[0], 1)[0] or name, pad(rows[1], 1)[0])
        out += [f"## {heading}", ""]
        if subtitle:
            out += [subtitle, ""]

        width = max(len(r) for r in rows[2:])
        head = pad(rows[2], width)
        body = rows[3:]

        if head[0] == "API ID":
            out += _api_endpoints(head, body)
        else:
            out += table(head, [pad(r, width) for r in body], width)
            out += [""]

    return "\n".join(out)


# ---------------------------------------------------------------- 실행


def newest(src: Path, needle: str) -> Path | None:
    """이름이 needle 로 시작하는 xlsx 중 가장 최근 것. 여러 개면 나머지도 알려준다."""
    hits = [p for p in src.glob("*.xlsx") if p.name.startswith(needle) and not p.name.startswith("~$")]
    if not hits:
        return None
    hits.sort(key=lambda p: p.stat().st_mtime, reverse=True)
    for other in hits[1:]:
        print(f"  건너뜀: {other.name}", file=sys.stderr)
    return hits[0]


def main() -> int:
    parser = argparse.ArgumentParser(description="구글 시트 xlsx를 docs/01_planning 마크다운으로 변환한다")
    parser.add_argument("--src", type=Path, default=Path.home() / "Downloads", help="xlsx를 받아둔 폴더")
    parser.add_argument("--out", type=Path, default=OUT_DIR, help="마크다운을 쓸 폴더")
    parser.add_argument("--only", choices=sorted(SOURCES), help="한 문서만 다시 뽑는다")
    args = parser.parse_args()

    today = dt.date.today().isoformat()
    builders = {"table": build_table_spec, "req": build_requirements, "api": build_api_spec}
    targets = [args.only] if args.only else list(SOURCES)

    missing: list[str] = []
    for key in targets:
        needle, out_name, sheet_name = SOURCES[key]
        path = newest(args.src, needle)
        if path is None:
            missing.append(f"  {sheet_name} — {args.src}에서 '{needle}'로 시작하는 xlsx를 못 찾음")
            continue

        workbook = load_workbook(path, data_only=True, read_only=True)
        text = builders[key](workbook, today).rstrip() + "\n"
        workbook.close()

        target = args.out / out_name
        before = target.read_text() if target.exists() else ""
        target.write_text(text)
        mark = "그대로" if before == text else f"{len(before)} → {len(text)}자"
        print(f"{out_name:<20} ← {path.name}  ({mark})")

    if missing:
        print("\n못 찾은 파일이 있습니다.", file=sys.stderr)
        print("\n".join(missing), file=sys.stderr)
        print("\n구글 시트에서 파일 → 다운로드 → Microsoft Excel(.xlsx)로 받아주세요.", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
