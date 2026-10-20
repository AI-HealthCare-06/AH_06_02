# /// script
# requires-python = ">=3.13"
# dependencies = ["openpyxl>=3.1"]
# ///
"""테이블 명세서의 v5 추가 테이블을 ERD·DDL·마이그레이션으로 뽑는다."""

import argparse
import re
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
TABLES = {"challenge_timer_sessions", "reward_draws", "adventurer_states"}
MARKER = "-- v5 API 추가 테이블 (구글 시트에서 생성)"


def read_groups(source: Path) -> dict[str, list[tuple]]:
    sheet = load_workbook(source, data_only=True)["테이블 명세"]
    groups: dict[str, list[tuple]] = {}
    current = ""
    for row in sheet.iter_rows(values_only=True):
        if row[0]:
            current = str(row[0])
        if current in TABLES and row[4]:
            groups.setdefault(current, []).append(tuple(row[4:10]))
    if set(groups) != TABLES:
        raise ValueError("명세서의 v5 테이블 3개를 확인해주세요.")
    return groups


def generate(source: Path, write_migration: bool = False) -> None:
    """ERD·DDL 사본을 다시 쓴다. 마이그레이션 8 은 이미 적용됐으므로 --migration 을 줄 때만 다시 쓴다."""
    groups = read_groups(source)
    column_count = sum(map(len, groups.values()))
    sql, diagram = [], []
    for table, fields in groups.items():
        columns, constraints, mermaid = [], [], []
        for name, kind, nullable, key, default, description in fields:
            # 기본값 0 도 값이다. None 일 때만 비우고, 시트가 숫자를 0.0 으로 주면 0 으로 적는다
            if isinstance(default, float) and default.is_integer():
                default = int(default)
            key, default = str(key or ""), "" if default is None else str(default)
            definition = f"  `{name}` {kind} " + ("NOT NULL" if nullable == "NN" else "NULL")
            if "PK" in key:
                definition += " AUTO_INCREMENT PRIMARY KEY"
            elif default and kind != "JSON":
                definition += f" DEFAULT {default}"
            if "UNIQUE" in key:
                definition += " UNIQUE"
            columns.append(definition)
            if "FK" in key:
                relation = re.search(r"→ (\w+)\.(\w+)", description)
                assert relation is not None
                parent, field = relation.groups()
                constraints.append(f"  FOREIGN KEY (`{name}`) REFERENCES `{parent}` (`{field}`)")
                left = "|o" if nullable == "Y" else "||"
                right = "o|" if "UNIQUE" in key else "o{"
                diagram.append(f'    {parent} {left}--{right} {table} : "{name}"')
            mermaid.append((f"        {str(kind).lower()} {name} " + key.replace("UNIQUE", "UK")).rstrip())
        sql.append(
            f"CREATE TABLE `{table}` (\n"
            + ",\n".join(columns + constraints)
            + ") ENGINE=InnoDB DEFAULT CHARSET=utf8mb4;"
        )
        diagram.append(f"    {table} {{\n" + "\n".join(mermaid) + "\n    }")
    ddl = "\n\n".join(sql)
    target = ROOT / "docs/01_planning/erd.sql"
    base = target.read_text().split(MARKER)[0].replace("SET FOREIGN_KEY_CHECKS = 1;", "").rstrip()
    target.write_text(base + "\n\n" + MARKER + "\n\n" + ddl + "\n\nSET FOREIGN_KEY_CHECKS = 1;\n")
    target = ROOT / "docs/01_planning/erd.md"
    base = target.read_text().split("\n## v5 API 추가 테이블")[0].rstrip()
    target.write_text(
        base + f"\n\n## v5 API 추가 테이블\n\n{date.today().isoformat()} 명세 사본에서 생성. 기존 테이블 유지, "
        f"{len(groups)}테이블·{column_count}컬럼 추가.\n\n```mermaid\nerDiagram\n" + "\n".join(diagram) + "\n```\n"
    )
    if not write_migration:
        return
    migration = ROOT / "app/core/db/migrations/models/8_20261018093000_v5_adventure.py"
    state = (
        migration.read_text().split("MODELS_STATE =", 1)[1]
        if migration.exists() and "MODELS_STATE =" in migration.read_text()
        else None
    )
    header = "from tortoise import BaseDBAsyncClient\n\nRUN_IN_TRANSACTION = True\n\n\n"
    upgrade = 'async def upgrade(db: BaseDBAsyncClient) -> str:\n    return """\n' + ddl + '\n"""\n'
    downgrade = (
        '\n\nasync def downgrade(db: BaseDBAsyncClient) -> str:\n    return """\n'
        "DROP TABLE adventurer_states; DROP TABLE reward_draws; DROP TABLE challenge_timer_sessions;\n"
        '\n"""\n'
    )
    migration.write_text(header + upgrade + downgrade)
    if state:
        with migration.open("a") as file:
            file.write("\n\nMODELS_STATE =" + state)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path)
    parser.add_argument("--migration", action="store_true", help="마이그레이션 8 파일도 다시 쓴다 (처음 만들 때만)")
    args = parser.parse_args()
    generate(args.source, args.migration)
